"""Extract PDF, DOCX and text evidence into page-linked, case-scoped chunks."""

import hashlib
from pathlib import Path
from zipfile import ZipFile

from docx import Document as WordDocument
from pypdf import PdfReader
from sqlalchemy import delete, select

from backend.config import settings
from database.models import Chunk, Task
from integrations.ai_provider import embed

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def extract_pages(path: Path):
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF: upload an unlocked copy")
        if len(reader.pages) > 500:
            raise ValueError("Split documents into at most 500 pages")
        pages = [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
        if any(len(text.strip()) < 15 for _, text in pages):
            raise ValueError(
                "One or more pages have no usable text. OCR the PDF and upload a searchable copy; no partial indexing was performed."
            )
        return pages
    if suffix == ".docx":
        with ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50_000_000:
                raise ValueError("Expanded DOCX exceeds 50 MB")
        doc = WordDocument(path)
        content = "\n".join(p.text for p in doc.paragraphs)
        content += "\n" + "\n".join(
            " | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows
        )
        return [(1, content)]  # DOCX has no stable page boundaries; UI labels these as sections.
    return [(1, path.read_text(encoding="utf-8-sig"))]


def run(db, case, document):
    path = Path(settings.storage_path) / document.storage_key
    if hashlib.sha256(path.read_bytes()).hexdigest() != document.sha256:
        raise ValueError("Stored document integrity check failed")
    pages = extract_pages(path)
    if not any(text.strip() for _, text in pages):
        raise ValueError("No readable text found")
    if sum(len(text) for _, text in pages) > 2_000_000:
        raise ValueError("Extracted content exceeds 2 million characters; split the document")
    db.execute(delete(Chunk).where(Chunk.document_id == document.id))
    chunks = []
    for page, text in pages:
        for start in range(0, len(text), 1200):
            content = text[start : start + 1500].strip()
            if content:
                chunks.append(
                    Chunk(
                        case_id=case.id,
                        document_id=document.id,
                        page=page,
                        text=content,
                    )
                )
    if settings.embeddings_enabled:
        for start in range(0, len(chunks), 64):
            batch = chunks[start : start + 64]
            vectors = embed([chunk.text for chunk in batch])
            for chunk, vector in zip(batch, vectors, strict=True):
                chunk.embedding = vector
    db.add_all(chunks)
    document.summary = f"{len(pages)} page/section(s), {len(chunks)} searchable passages. Category supplied by uploader; content requires review."
    document.status, document.error = "ready", ""
    collection = db.scalar(
        select(Task).where(Task.dedupe_key == f"{case.id}:collect:{document.category}")
    )
    if collection:
        collection.completed = True
